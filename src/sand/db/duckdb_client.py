"""DuckDB database client."""

from __future__ import annotations

import json
import re
import threading
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd

METADATA_TABLE = "_sand_metadata"


class DuckDBClient:
    """DuckDB wrapper used by ingest, queries, and chat."""

    def __init__(self, path: str | Path, *, read_only: bool = False, owns_connection: bool = True):
        self.path = Path(path)
        self.read_only = read_only
        self.owns_connection = owns_connection
        if not read_only:
            self.path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        try:
            self._conn = duckdb.connect(str(self.path), read_only=read_only)
        except duckdb.IOException:
            raise
        self._apply_timeout()
        if not read_only:
            self._ensure_metadata()

    def _ensure_metadata(self) -> None:
        # Use raw conn so timeout wrapper isn't required during init
        self._conn.execute(
            f"""
            CREATE TABLE IF NOT EXISTS {METADATA_TABLE} (
                table_name VARCHAR PRIMARY KEY,
                source_file VARCHAR,
                sheet_name VARCHAR,
                original_columns VARCHAR,
                row_count BIGINT,
                created_at TIMESTAMP DEFAULT now()
            )
            """
        )

    def _apply_timeout(self) -> None:
        try:
            from sand.core.config import get_settings

            seconds = float(get_settings().query_timeout_seconds)
            if seconds > 0:
                # DuckDB interrupt is driven from Python; store for execute wrappers
                self._timeout_seconds = seconds
            else:
                self._timeout_seconds = 0.0
        except Exception:
            self._timeout_seconds = 30.0

    def _run(self, fn):  # noqa: ANN001
        timeout = getattr(self, "_timeout_seconds", 0.0) or 0.0
        if timeout <= 0:
            return fn()
        timer = threading.Timer(timeout, self._interrupt)
        timer.daemon = True
        timer.start()
        try:
            return fn()
        except Exception as exc:
            msg = str(exc).lower()
            if "interrupt" in msg or "canceled" in msg or "cancelled" in msg:
                raise TimeoutError(f"Query exceeded SAND_QUERY_TIMEOUT_SECONDS={timeout}") from exc
            raise
        finally:
            timer.cancel()

    def _interrupt(self) -> None:
        try:
            self._conn.interrupt()
        except Exception:
            pass

    def execute(self, sql: str, params: tuple[Any, ...] | list[Any] | None = None) -> None:
        with self._lock:
            def _do() -> None:
                if params:
                    self._conn.execute(sql, list(params))
                else:
                    self._conn.execute(sql)

            self._run(_do)

    def executemany(self, sql: str, params_seq: list[tuple[Any, ...]]) -> None:
        with self._lock:
            self._run(lambda: self._conn.executemany(sql, params_seq))

    def fetchall(self, sql: str, params: tuple[Any, ...] | list[Any] | None = None) -> list[tuple[Any, ...]]:
        with self._lock:
            def _do() -> list[tuple[Any, ...]]:
                if params:
                    cur = self._conn.execute(sql, list(params))
                else:
                    cur = self._conn.execute(sql)
                return [tuple(row) for row in cur.fetchall()]

            return self._run(_do)

    def to_dataframe(self, sql: str, params: tuple[Any, ...] | list[Any] | None = None) -> pd.DataFrame:
        with self._lock:
            def _do() -> pd.DataFrame:
                if params:
                    return self._conn.execute(sql, list(params)).df()
                return self._conn.execute(sql).df()

            return self._run(_do)

    def copy_to_csv(self, sql: str, dest: Path) -> None:
        """Stream a SELECT to CSV on disk via DuckDB COPY (avoids full pandas buffer)."""
        path_lit = str(Path(dest).resolve()).replace("'", "''")
        cleaned = sql.strip().rstrip(";")
        self.execute(f"COPY ({cleaned}) TO '{path_lit}' (HEADER, DELIMITER ',', FORMAT CSV)")

    def copy_to_xlsx(self, sql: str, dest: Path) -> None:
        """Export a SELECT to XLSX via CSV staging + openpyxl write-only (bounded memory)."""
        import csv
        import tempfile

        from openpyxl import Workbook

        tmp = tempfile.NamedTemporaryFile(prefix="sand_xlsx_", suffix=".csv", delete=False)
        tmp_csv = Path(tmp.name)
        tmp.close()
        try:
            self.copy_to_csv(sql, tmp_csv)
            wb = Workbook(write_only=True)
            ws = wb.create_sheet(title="export")
            with tmp_csv.open(newline="", encoding="utf-8") as fh:
                for row in csv.reader(fh):
                    ws.append(row)
            dest = Path(dest)
            dest.parent.mkdir(parents=True, exist_ok=True)
            wb.save(dest)
        finally:
            tmp_csv.unlink(missing_ok=True)

    def table_names(self) -> list[str]:
        rows = self.fetchall(
            """
            SELECT table_name
            FROM information_schema.tables
            WHERE table_schema = 'main' AND table_type = 'BASE TABLE'
            ORDER BY table_name
            """
        )
        return [r[0] for r in rows if not str(r[0]).startswith("_sand")]

    def schema(self) -> dict[str, list[dict[str, str]]]:
        result: dict[str, list[dict[str, str]]] = {}
        for name in self.table_names():
            cols = self.fetchall(
                """
                SELECT column_name, data_type
                FROM information_schema.columns
                WHERE table_schema = 'main' AND table_name = ?
                ORDER BY ordinal_position
                """,
                (name,),
            )
            result[name] = [{"name": c[0], "type": c[1] or "VARCHAR"} for c in cols]
        return result

    def write_dataframe(self, df: pd.DataFrame, table_name: str, if_exists: str = "replace") -> int:
        if self.read_only:
            raise RuntimeError("Cannot write to a read-only DuckDB connection")
        safe = sanitize_table_name(table_name)
        with self._lock:
            rel_name = "_sand_incoming_df"
            self._conn.register(rel_name, df)
            try:
                if if_exists == "replace":
                    self._conn.execute(f"CREATE OR REPLACE TABLE {_quote_ident(safe)} AS SELECT * FROM {rel_name}")
                elif if_exists == "fail":
                    exists = self._conn.execute(
                        """
                        SELECT 1 FROM information_schema.tables
                        WHERE table_schema = 'main' AND table_name = ?
                        """,
                        [safe],
                    ).fetchone()
                    if exists:
                        raise ValueError(f"Table '{safe}' already exists")
                    self._conn.execute(f"CREATE TABLE {_quote_ident(safe)} AS SELECT * FROM {rel_name}")
                else:
                    raise ValueError("if_exists must be 'replace' or 'fail'")
            finally:
                try:
                    self._conn.unregister(rel_name)
                except Exception:
                    pass
        return len(df)

    def create_table_as(self, table_name: str, sql: str) -> int:
        """Materialize a SELECT into a table without pulling rows into Python."""
        if self.read_only:
            raise RuntimeError("Cannot write to a read-only DuckDB connection")
        safe = sanitize_table_name(table_name)
        self.execute(f"CREATE OR REPLACE TABLE {_quote_ident(safe)} AS {sql}")
        return int(self.fetchall(f"SELECT COUNT(*) FROM {_quote_ident(safe)}")[0][0])

    def checkpoint(self) -> None:
        """Flush WAL into the main database file for safe copy/export."""
        if self.read_only:
            return
        try:
            self.execute("CHECKPOINT")
        except Exception:
            # Older / busy connections may reject; best-effort
            pass

    def _ensure_excel(self) -> None:
        with self._lock:
            if getattr(self, "_excel_ready", False):
                return
            self._conn.execute("INSTALL excel")
            self._conn.execute("LOAD excel")
            self._excel_ready = True

    def _ctas_from_select(self, table_name: str, select_sql: str, *, if_exists: str) -> tuple[int, list[str]]:
        safe = sanitize_table_name(table_name)
        if if_exists == "replace":
            self.execute(f"CREATE OR REPLACE TABLE {_quote_ident(safe)} AS {select_sql}")
        elif if_exists == "fail":
            if safe in self.table_names():
                raise ValueError(f"Table '{safe}' already exists")
            self.execute(f"CREATE TABLE {_quote_ident(safe)} AS {select_sql}")
        else:
            raise ValueError("if_exists must be 'replace' or 'fail'")
        cols = [c["name"] for c in self.schema().get(safe, [])]
        count = int(self.fetchall(f"SELECT COUNT(*) FROM {_quote_ident(safe)}")[0][0])
        return count, cols

    def ingest_csv(self, path: Path, table_name: str, *, if_exists: str = "replace") -> tuple[int, list[str]]:
        """Load a CSV via DuckDB's native reader (avoids full pandas materialization)."""
        if self.read_only:
            raise RuntimeError("Cannot write to a read-only DuckDB connection")
        path_lit = str(Path(path).resolve()).replace("'", "''")
        select_sql = f"SELECT * FROM read_csv_auto('{path_lit}', header=true, sample_size=-1)"
        return self._ctas_from_select(table_name, select_sql, if_exists=if_exists)

    def ingest_parquet(self, path: Path, table_name: str, *, if_exists: str = "replace") -> tuple[int, list[str]]:
        """Load a Parquet file via DuckDB's native reader."""
        if self.read_only:
            raise RuntimeError("Cannot write to a read-only DuckDB connection")
        path_lit = str(Path(path).resolve()).replace("'", "''")
        select_sql = f"SELECT * FROM read_parquet('{path_lit}')"
        return self._ctas_from_select(table_name, select_sql, if_exists=if_exists)

    def ingest_xlsx_sheet(
        self,
        path: Path,
        table_name: str,
        *,
        sheet: str,
        if_exists: str = "replace",
    ) -> tuple[int, list[str]]:
        """Load one Excel sheet via DuckDB ``read_xlsx`` (excel extension)."""
        if self.read_only:
            raise RuntimeError("Cannot write to a read-only DuckDB connection")
        self._ensure_excel()
        path_lit = str(Path(path).resolve()).replace("'", "''")
        sheet_lit = sheet.replace("'", "''")
        select_sql = (
            f"SELECT * FROM read_xlsx('{path_lit}', sheet='{sheet_lit}', header=true, all_varchar=false)"
        )
        return self._ctas_from_select(table_name, select_sql, if_exists=if_exists)

    def register_table(
        self,
        table_name: str,
        *,
        source_file: str,
        sheet_name: str,
        original_columns: list[str],
        row_count: int,
    ) -> None:
        self.execute(
            f"""
            INSERT INTO {METADATA_TABLE}
                (table_name, source_file, sheet_name, original_columns, row_count)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT (table_name) DO UPDATE SET
                source_file = excluded.source_file,
                sheet_name = excluded.sheet_name,
                original_columns = excluded.original_columns,
                row_count = excluded.row_count,
                created_at = now()
            """,
            (table_name, source_file, sheet_name, json.dumps(original_columns), row_count),
        )

    def metadata(self) -> pd.DataFrame:
        return self.to_dataframe(f"SELECT * FROM {METADATA_TABLE}")

    def close(self) -> None:
        if not self.owns_connection:
            return
        with self._lock:
            try:
                self._conn.close()
            except Exception:
                pass

    def __enter__(self) -> DuckDBClient:
        return self

    def __exit__(self, *args: object) -> None:
        self.close()


def sanitize_table_name(name: str) -> str:
    cleaned = re.sub(r"[^0-9a-zA-Z_]+", "_", name.strip())
    cleaned = re.sub(r"_+", "_", cleaned).strip("_").lower()
    if not cleaned:
        cleaned = "sheet"
    if cleaned[0].isdigit():
        cleaned = f"t_{cleaned}"
    return cleaned[:64]


def _quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'
